"""`cli.run` end to end (CODING_STANDARDS.md section 5): a real `.xlsx` fixture and a
real mapping file, run over the `conftest` engine and movable clock, with the business
opened through the `client` fixture and every figure read back through the HTTP API.
Never by building rows by hand (the plan's seam for this subtask).
"""

from __future__ import annotations

import io
import logging
import re
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.engine import Engine
from sqlalchemy.orm import Session

from app.accounts.adapters.tables import BusinessRow
from app.customers.adapters.import_port import CustomersImport
from app.customers.adapters.tables import CustomerRow
from app.db import make_engine
from app.main import create_app
from app.migrations.runner import upgrade_to_head
from app.sales.adapters.import_port import SalesImport
from app.settings import Settings
from app.shared.copy import OperatorText
from app.shared.transaction import Transaction
from app.spreadsheet_import.adapters import layout
from app.spreadsheet_import.adapters.cli import main, run
from tests.conftest import TEST_HASHER, MovableClock, open_business, signed_in
from tests.fakes.telegram import FakeTelegram

# No report line, exception message or log line ever carries a phone number, not even
# a run of 4 or more of its digits (the plan's invariant).
_DIGIT_RUN = re.compile(r"\d{4,}")

VENTAS_TITLES = (
    "Fecha",
    "Variedad",
    "Gramos",
    "Pago",
    "Total",
    "Cliente",
    "Tipo",
    "Envío cobrado",
    "Precio especial (total)",
    "Nota",
)
ENTRADAS_TITLES = ("Fecha", "Variedad", "Gramos", "Origen", "Costo total", "Nota", "Revisar")
GASTOS_TITLES = ("Fecha", "Categoría", "Monto", "Detalle", "Revisar")

VARIETY = "Jack Herer"
CUSTOMER = "Juan Perez"
PHONE = "1123456789"
SALE_TOTAL = 10_000


def _titles(ws: Worksheet, titles: tuple[str, ...]) -> None:
    for index, title in enumerate(titles, start=1):
        ws.cell(row=layout.TITLE_ROW, column=index, value=title)


def _workbook(
    *,
    sale_date: date = date(2026, 1, 10),
    payment: str = "Pendiente",
    threshold: float | None = 5,
) -> Workbook:
    workbook = Workbook()
    active = workbook.active
    assert active is not None
    workbook.remove(active)

    ajustes = workbook.create_sheet(layout.AJUSTES_SHEET)
    ajustes.cell(row=7, column=1, value=VARIETY)
    if threshold is not None:
        ajustes.cell(row=7, column=6, value=threshold)
    ajustes.cell(row=21, column=1, value="Insumos")
    ajustes.cell(row=37, column=1, value=VARIETY)
    ajustes.cell(row=37, column=2, value=date(2026, 1, 1))
    ajustes.cell(row=37, column=3, value=1000)

    catalogo = workbook.create_sheet(layout.CATALOGO_SHEET)
    catalogo.cell(row=6, column=1, value=VARIETY)
    catalogo.cell(row=6, column=2, value="Sí")
    catalogo.cell(row=18, column=2, value="Catálogo de la casa")
    catalogo.cell(row=19, column=2, value="Gracias por su compra")

    ventas = workbook.create_sheet(layout.VENTAS_SHEET)
    _titles(ventas, VENTAS_TITLES)
    ventas.cell(row=2, column=1, value=sale_date)
    ventas.cell(row=2, column=2, value=VARIETY)
    ventas.cell(row=2, column=3, value=10)
    ventas.cell(row=2, column=4, value=payment)
    ventas.cell(row=2, column=5, value=SALE_TOTAL)
    ventas.cell(row=2, column=6, value=CUSTOMER)
    ventas.cell(row=2, column=7, value="")

    entradas = workbook.create_sheet(layout.ENTRADAS_SHEET)
    _titles(entradas, ENTRADAS_TITLES)
    entradas.cell(row=2, column=1, value=date(2026, 1, 1))
    entradas.cell(row=2, column=2, value=VARIETY)
    entradas.cell(row=2, column=3, value=50)
    entradas.cell(row=2, column=4, value="Compra")
    entradas.cell(row=2, column=5, value=20_000)

    gastos = workbook.create_sheet(layout.GASTOS_SHEET)
    _titles(gastos, GASTOS_TITLES)
    gastos.cell(row=2, column=1, value=date(2026, 1, 2))
    gastos.cell(row=2, column=2, value="Insumos")
    gastos.cell(row=2, column=3, value=500)
    gastos.cell(row=2, column=4, value="Bolsas")

    return workbook


def _save_workbook(workbook: Workbook, tmp_path: Path, name: str = "sheet.xlsx") -> Path:
    path = tmp_path / name
    workbook.save(path)
    return path


def _save_mapping(
    tmp_path: Path, rows: tuple[tuple[str, str], ...], name: str = "mapping.csv"
) -> Path:
    path = tmp_path / name
    with path.open("w", encoding="utf-8", newline="") as handle:
        handle.write("name,phone\n")
        for row_name, phone in rows:
            handle.write(f"{row_name},{phone}\n")
    return path


def _business_id(engine: Engine, name: str = "Primer Negocio") -> int:
    with Session(engine) as session:
        row = session.scalars(select(BusinessRow).where(BusinessRow.name == name)).one()
        assert row.id is not None
        return row.id


def test_dry_run_writes_nothing_and_reports_the_counts(
    client: TestClient,
    engine: Engine,
    clock: MovableClock,
    tmp_path: Path,
    operator_text_en: OperatorText,
) -> None:
    business_id = _business_id(engine)
    xlsx = _save_workbook(_workbook(), tmp_path)
    mapping = _save_mapping(tmp_path, ((CUSTOMER, PHONE),))
    out, err = io.StringIO(), io.StringIO()

    code = run(
        [str(business_id), str(xlsx), "--mapping", str(mapping), "--dry-run"],
        engine=engine,
        clock=clock,
        out=out,
        err=err,
        text=operator_text_en,
    )

    assert code == 0
    assert err.getvalue() == ""
    assert "Sales: 1" in out.getvalue()
    assert "Writable: yes" in out.getvalue()
    assert client.get("/api/products").json() == []
    assert client.get("/customers").json() == []
    assert client.get("/stock").json()["products"] == []


def test_real_run_imports_and_the_figures_match(
    client: TestClient,
    engine: Engine,
    clock: MovableClock,
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
    operator_text_en: OperatorText,
) -> None:
    business_id = _business_id(engine)
    xlsx = _save_workbook(_workbook(), tmp_path)
    mapping = _save_mapping(tmp_path, ((CUSTOMER, PHONE),))
    out, err = io.StringIO(), io.StringIO()

    with caplog.at_level(logging.DEBUG):
        code = run(
            [str(business_id), str(xlsx), "--mapping", str(mapping)],
            engine=engine,
            clock=clock,
            out=out,
            err=err,
            text=operator_text_en,
        )

    assert code == 0, err.getvalue()
    assert err.getvalue() == ""
    assert caplog.records == []
    assert PHONE not in out.getvalue()
    assert not _DIGIT_RUN.search(out.getvalue())

    varieties = client.get("/api/products").json()
    assert len(varieties) == 1
    assert varieties[0]["name"] == VARIETY
    assert varieties[0]["current_price"]["list_price"] == 1000

    catalog = client.get("/catalog").json()
    assert catalog["header"] == "Catálogo de la casa"
    assert catalog["footer"] == "Gracias por su compra"

    stock = client.get("/stock").json()
    assert len(stock["products"]) == 1
    # 50 g bought in, 10 g sold: 40 g left.
    assert stock["products"][0]["quantity"] == 40.0
    assert stock["products"][0]["threshold"] == 5.0

    customers = client.get("/customers").json()
    assert len(customers) == 1
    assert customers[0]["name"] == CUSTOMER
    assert customers[0]["balance"] == SALE_TOTAL
    assert customers[0]["owing_since"] == "2026-01-10"


def _set_zone(engine: Engine, business_id: int, zone: str) -> None:
    with Session(engine) as session:
        session.get_one(BusinessRow, business_id).time_zone = zone
        session.commit()


def _dry_run(
    client: TestClient,
    engine: Engine,
    clock: MovableClock,
    tmp_path: Path,
    operator_text_en: OperatorText,
) -> tuple[int, str, str]:
    xlsx = _save_workbook(_workbook(), tmp_path)
    mapping = _save_mapping(tmp_path, ((CUSTOMER, PHONE),))
    out, err = io.StringIO(), io.StringIO()
    code = run(
        [str(_business_id(engine)), str(xlsx), "--mapping", str(mapping), "--dry-run"],
        engine=engine,
        clock=clock,
        out=out,
        err=err,
        text=operator_text_en,
    )
    return code, out.getvalue(), err.getvalue()


def test_the_import_counts_today_in_the_target_businesss_zone(
    client: TestClient,
    engine: Engine,
    clock: MovableClock,
    tmp_path: Path,
    operator_text_en: OperatorText,
) -> None:
    # The sheet's latest sale is dated 10 January 2026. At this instant it is still the
    # 9th in Buenos Aires and already the 10th on Kiritimati (UTC+14).
    clock.set(datetime(2026, 1, 9, 20, 0, tzinfo=UTC))

    code, out, _ = _dry_run(client, engine, clock, tmp_path, operator_text_en)
    assert code == 0
    assert "Writable: no" in out

    _set_zone(engine, _business_id(engine), "Pacific/Kiritimati")
    code, out, _ = _dry_run(client, engine, clock, tmp_path, operator_text_en)
    assert code == 0
    assert "Writable: yes" in out


def test_a_target_whose_zone_cannot_be_loaded_stops_the_import_naming_no_data(
    client: TestClient,
    engine: Engine,
    clock: MovableClock,
    tmp_path: Path,
    operator_text_en: OperatorText,
) -> None:
    _set_zone(engine, _business_id(engine), "Not/AZone")

    code, out, err = _dry_run(client, engine, clock, tmp_path, operator_text_en)

    assert code == 1
    assert out == ""
    assert err == "refused: import.time_zone_unavailable\n"


def test_an_unknown_business_is_refused_with_no_echo(
    engine: Engine, clock: MovableClock, tmp_path: Path, operator_text_en: OperatorText
) -> None:
    upgrade_to_head(engine)
    xlsx = _save_workbook(_workbook(), tmp_path)
    mapping = _save_mapping(tmp_path, ((CUSTOMER, PHONE),))
    out, err = io.StringIO(), io.StringIO()

    code = run(
        ["999999", str(xlsx), "--mapping", str(mapping)],
        engine=engine,
        clock=clock,
        out=out,
        err=err,
        text=operator_text_en,
    )

    assert code == 1
    assert err.getvalue() == "refused: import.business_unknown\n"
    assert out.getvalue() == ""


def test_a_business_with_records_is_refused_with_nothing_written(
    client: TestClient,
    engine: Engine,
    clock: MovableClock,
    tmp_path: Path,
    operator_text_en: OperatorText,
) -> None:
    client.post("/api/products", json={"name": "Existente"})
    business_id = _business_id(engine)
    xlsx = _save_workbook(_workbook(), tmp_path)
    mapping = _save_mapping(tmp_path, ((CUSTOMER, PHONE),))
    out, err = io.StringIO(), io.StringIO()

    code = run(
        [str(business_id), str(xlsx), "--mapping", str(mapping)],
        engine=engine,
        clock=clock,
        out=out,
        err=err,
        text=operator_text_en,
    )

    assert code == 1
    assert err.getvalue() == ""
    assert "import.business_has_records" in out.getvalue()
    varieties = client.get("/api/products").json()
    assert [v["name"] for v in varieties] == ["Existente"]


def test_imported_records_are_invisible_from_another_business(
    client: TestClient,
    other_client: TestClient,
    engine: Engine,
    clock: MovableClock,
    tmp_path: Path,
    operator_text_en: OperatorText,
) -> None:
    business_id = _business_id(engine)
    xlsx = _save_workbook(_workbook(), tmp_path)
    mapping = _save_mapping(tmp_path, ((CUSTOMER, PHONE),))
    out, err = io.StringIO(), io.StringIO()

    code = run(
        [str(business_id), str(xlsx), "--mapping", str(mapping)],
        engine=engine,
        clock=clock,
        out=out,
        err=err,
        text=operator_text_en,
    )

    assert code == 0, err.getvalue()
    assert other_client.get("/api/products").json() == []
    assert other_client.get("/customers").json() == []
    assert other_client.get("/stock").json()["products"] == []


def test_a_malformed_mapping_on_a_real_run_is_refused_with_no_echo(
    client: TestClient,
    engine: Engine,
    clock: MovableClock,
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
    operator_text_en: OperatorText,
) -> None:
    business_id = _business_id(engine)
    xlsx = _save_workbook(_workbook(), tmp_path)
    mapping_path = tmp_path / "mapping.csv"
    # A header that doesn't match at all: every entry is unreachable, so the
    # problem is located by line number alone, never by content.
    mapping_path.write_text(f"nombre,telefono\n{CUSTOMER},{PHONE}\n", encoding="utf-8")
    out, err = io.StringIO(), io.StringIO()

    with caplog.at_level(logging.DEBUG):
        code = run(
            [str(business_id), str(xlsx), "--mapping", str(mapping_path)],
            engine=engine,
            clock=clock,
            out=out,
            err=err,
            text=operator_text_en,
        )

    assert code == 1
    report = out.getvalue()
    assert PHONE not in report
    assert CUSTOMER not in report
    assert not _DIGIT_RUN.search(report)
    assert not _DIGIT_RUN.search(err.getvalue())
    assert caplog.records == []
    assert client.get("/customers").json() == []


def test_a_mapping_entry_with_a_blank_phone_imports_the_customer_without_one(
    client: TestClient,
    engine: Engine,
    clock: MovableClock,
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
    operator_text_en: OperatorText,
) -> None:
    # A blank phone cell means no number (ADR-0021): it is not a problem, and the
    # customer is imported under their name alone.
    business_id = _business_id(engine)
    xlsx = _save_workbook(_workbook(), tmp_path)
    mapping_path = tmp_path / "mapping.csv"
    mapping_path.write_text(f"name,phone\n{CUSTOMER},\n", encoding="utf-8")
    out, err = io.StringIO(), io.StringIO()

    with caplog.at_level(logging.DEBUG):
        code = run(
            [str(business_id), str(xlsx), "--mapping", str(mapping_path)],
            engine=engine,
            clock=clock,
            out=out,
            err=err,
            text=operator_text_en,
        )

    assert code == 0, out.getvalue()
    report = out.getvalue()
    assert "Customers without a phone number: 1" in report
    assert not _DIGIT_RUN.search(report)
    assert not _DIGIT_RUN.search(err.getvalue())
    assert caplog.records == []
    customers = client.get("/customers").json()
    assert len(customers) == 1
    assert customers[0]["phone"] is None


def test_no_mapping_flag_imports_every_customer_without_a_phone(
    client: TestClient,
    engine: Engine,
    clock: MovableClock,
    tmp_path: Path,
    operator_text_en: OperatorText,
) -> None:
    business_id = _business_id(engine)
    xlsx = _save_workbook(_workbook(), tmp_path)
    out, err = io.StringIO(), io.StringIO()

    code = run(
        [str(business_id), str(xlsx)],
        engine=engine,
        clock=clock,
        out=out,
        err=err,
        text=operator_text_en,
    )

    assert code == 0, err.getvalue()
    report = out.getvalue()
    assert "Customers without a phone number: 1" in report
    customers = client.get("/customers").json()
    assert len(customers) == 1
    assert customers[0]["name"] == CUSTOMER
    assert customers[0]["phone"] is None


def test_a_partial_mapping_imports_unmapped_names_without_a_phone(
    client: TestClient,
    engine: Engine,
    clock: MovableClock,
    tmp_path: Path,
    operator_text_en: OperatorText,
) -> None:
    workbook = _workbook()
    ventas = workbook[layout.VENTAS_SHEET]
    ventas.cell(row=3, column=1, value=date(2026, 1, 11))
    ventas.cell(row=3, column=2, value=VARIETY)
    ventas.cell(row=3, column=3, value=5)
    ventas.cell(row=3, column=4, value="Pendiente")
    ventas.cell(row=3, column=5, value=5000)
    ventas.cell(row=3, column=6, value=CUSTOMER2)
    ventas.cell(row=3, column=7, value="")
    xlsx = _save_workbook(workbook, tmp_path)
    mapping = _save_mapping(tmp_path, ((CUSTOMER, PHONE),))
    business_id = _business_id(engine)
    out, err = io.StringIO(), io.StringIO()

    code = run(
        [str(business_id), str(xlsx), "--mapping", str(mapping)],
        engine=engine,
        clock=clock,
        out=out,
        err=err,
        text=operator_text_en,
    )

    assert code == 0, err.getvalue()
    report = out.getvalue()
    assert "Customers without a phone number: 1" in report
    customers = {c["name"]: c for c in client.get("/customers").json()}
    assert customers[CUSTOMER]["phone"] is not None
    assert customers[CUSTOMER2]["phone"] is None


def test_a_well_formed_mapping_row_with_a_malformed_phone_blocks_with_no_echo(
    client: TestClient,
    engine: Engine,
    clock: MovableClock,
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
    operator_text_en: OperatorText,
) -> None:
    # The row itself is well formed (one name, one non-blank phone cell): the
    # phone text just doesn't normalise. check_customer's own phone_code must
    # still block the run, filed at the mapping line, never at the Ventas row.
    business_id = _business_id(engine)
    xlsx = _save_workbook(_workbook(), tmp_path)
    mapping = _save_mapping(tmp_path, ((CUSTOMER, "abc"),))
    out, err = io.StringIO(), io.StringIO()

    with caplog.at_level(logging.DEBUG):
        code = run(
            [str(business_id), str(xlsx), "--mapping", str(mapping)],
            engine=engine,
            clock=clock,
            out=out,
            err=err,
            text=operator_text_en,
        )

    assert code == 1
    report = out.getvalue()
    assert "customer.phone_invalid" in report
    assert "Mapping row 2 [phone]: customer.phone_invalid" in report
    assert "Cliente]" not in report
    assert CUSTOMER not in report
    # The malformed phone itself ("abc") has no digit run, so _DIGIT_RUN alone
    # could not catch an echo of it: the literal text is checked too (the
    # mapping's phone value must never reach the report, digits or not).
    assert "abc" not in report
    assert not _DIGIT_RUN.search(report)
    assert not _DIGIT_RUN.search(err.getvalue())
    assert caplog.records == []
    assert client.get("/customers").json() == []
    assert client.get("/api/products").json() == []
    assert client.get("/stock").json()["products"] == []


def test_the_dry_run_count_of_customers_without_a_phone_matches_the_real_run(
    client: TestClient,
    engine: Engine,
    clock: MovableClock,
    tmp_path: Path,
    operator_text_en: OperatorText,
) -> None:
    workbook = _workbook()
    ventas = workbook[layout.VENTAS_SHEET]
    ventas.cell(row=3, column=1, value=date(2026, 1, 11))
    ventas.cell(row=3, column=2, value=VARIETY)
    ventas.cell(row=3, column=3, value=5)
    ventas.cell(row=3, column=4, value="Pendiente")
    ventas.cell(row=3, column=5, value=5000)
    ventas.cell(row=3, column=6, value=CUSTOMER2)
    ventas.cell(row=3, column=7, value="")
    xlsx = _save_workbook(workbook, tmp_path)
    mapping = _save_mapping(tmp_path, ((CUSTOMER, PHONE),))
    business_id = _business_id(engine)

    dry_out, dry_err = io.StringIO(), io.StringIO()
    dry_code = run(
        [str(business_id), str(xlsx), "--mapping", str(mapping), "--dry-run"],
        engine=engine,
        clock=clock,
        out=dry_out,
        err=dry_err,
        text=operator_text_en,
    )

    assert dry_code == 0, dry_err.getvalue()
    assert "Customers without a phone number: 1" in dry_out.getvalue()
    assert client.get("/customers").json() == []

    real_out, real_err = io.StringIO(), io.StringIO()
    real_code = run(
        [str(business_id), str(xlsx), "--mapping", str(mapping)],
        engine=engine,
        clock=clock,
        out=real_out,
        err=real_err,
        text=operator_text_en,
    )

    assert real_code == 0, real_err.getvalue()
    assert "Customers without a phone number: 1" in real_out.getvalue()
    customers = client.get("/customers").json()
    without_phone = [c for c in customers if c["phone"] is None]
    assert len(without_phone) == 1


def test_the_last_writer_forced_to_raise_leaves_nothing_written(
    client: TestClient,
    engine: Engine,
    clock: MovableClock,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    operator_text_en: OperatorText,
) -> None:
    # Ventas is the last kind written (varieties -> ... -> Entradas -> Gastos ->
    # Ventas): forcing the real `SalesImport.write_sale` port method to raise proves
    # everything staged before it (varieties, stock, catalog, categories, customers)
    # is rolled back too, not just skipped from this point on.
    def _boom(*_args: Any, **_kwargs: Any) -> int:
        raise RuntimeError(f"forced failure, phone {PHONE}")

    monkeypatch.setattr(SalesImport, "write_sale", _boom)

    business_id = _business_id(engine)
    xlsx = _save_workbook(_workbook(), tmp_path)
    mapping = _save_mapping(tmp_path, ((CUSTOMER, PHONE),))
    out, err = io.StringIO(), io.StringIO()

    code = run(
        [str(business_id), str(xlsx), "--mapping", str(mapping)],
        engine=engine,
        clock=clock,
        out=out,
        err=err,
        text=operator_text_en,
    )

    assert code == 1
    assert err.getvalue() == "refused: RuntimeError\n"
    assert PHONE not in out.getvalue()
    assert PHONE not in err.getvalue()
    assert not _DIGIT_RUN.search(out.getvalue())
    assert not _DIGIT_RUN.search(err.getvalue())
    assert client.get("/api/products").json() == []
    assert client.get("/customers").json() == []
    assert client.get("/stock").json()["products"] == []
    # Only the 4 seeded categories remain: the fixture's own "Insumos" never landed.
    assert len(client.get("/expense-categories").json()) == 4
    assert client.get("/catalog").json()["header"] != "Catálogo de la casa"


def test_a_dry_run_and_a_refused_real_run_over_a_blocking_fixture_write_nothing(
    client: TestClient,
    engine: Engine,
    clock: MovableClock,
    tmp_path: Path,
    operator_text_en: OperatorText,
) -> None:
    # An unknown Entradas origin blocks the whole plan; nothing else about the
    # fixture is invalid, so every other row would otherwise be written fine.
    workbook = _workbook()
    workbook[layout.ENTRADAS_SHEET].cell(row=2, column=4, value="Robo")
    xlsx = _save_workbook(workbook, tmp_path)
    mapping = _save_mapping(tmp_path, ((CUSTOMER, PHONE),))
    business_id = _business_id(engine)

    for argv_tail in (["--dry-run"], []):
        out, err = io.StringIO(), io.StringIO()
        code = run(
            [str(business_id), str(xlsx), "--mapping", str(mapping), *argv_tail],
            engine=engine,
            clock=clock,
            out=out,
            err=err,
            text=operator_text_en,
        )
        assert code in (0, 1)
        assert err.getvalue() == ""
        assert "Writable: no" in out.getvalue()
        assert "import.origin_unknown" in out.getvalue()
        assert client.get("/api/products").json() == []
        assert client.get("/customers").json() == []
        assert client.get("/stock").json()["products"] == []
        assert client.get("/catalog").json()["header"] != "Catálogo de la casa"


def test_a_real_duplicate_phone_is_refused_by_its_own_clean_code(
    client: TestClient,
    engine: Engine,
    clock: MovableClock,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    operator_text_en: OperatorText,
) -> None:
    """A second writer lands a customer with the same phone number into this business
    *during* the write — the known, accepted race of concurrent writes (ADR-0018's
    "Known risks") — rather than before it starts, which eligibility would simply
    refuse up front instead. The real `CustomersImport.write_customer` is made to
    insert that racing row, through a second session, right before it does its own
    real work. The collision is real: SQLite's own unique constraint on
    `(business_id, phone)` fires on flush, and `SqlCustomers`
    (app/customers/adapters/sql_customers.py) always turns that `IntegrityError` into
    the clean `PhoneTaken` before it leaves the adapter, so what reaches `cli.run`
    already carries only a fixed code, never the driver's own message — which, with a
    real phone number bound, is exactly what `hide_parameters=True` on `make_engine`
    would have kept out of a log line too, had this been logged."""
    business_id = _business_id(engine)
    original = CustomersImport.write_customer
    raced = False

    def _racing_write_customer(
        self: CustomersImport, name: str, phone_key: str | None, vip: bool, transaction: Transaction
    ) -> int:
        nonlocal raced
        if not raced:
            raced = True
            with Session(engine) as racer:
                racer.add(CustomerRow(business_id=business_id, name="Otro", phone=phone_key))
                racer.commit()
        return original(self, name, phone_key, vip, transaction)

    monkeypatch.setattr(CustomersImport, "write_customer", _racing_write_customer)
    xlsx = _save_workbook(_workbook(), tmp_path)
    mapping = _save_mapping(tmp_path, ((CUSTOMER, PHONE),))
    out, err = io.StringIO(), io.StringIO()

    code = run(
        [str(business_id), str(xlsx), "--mapping", str(mapping)],
        engine=engine,
        clock=clock,
        out=out,
        err=err,
        text=operator_text_en,
    )

    assert code == 1
    assert err.getvalue() == "refused: customer.phone_taken\n"
    assert PHONE not in out.getvalue()
    assert CUSTOMER not in out.getvalue()
    assert not _DIGIT_RUN.search(out.getvalue())
    assert not _DIGIT_RUN.search(err.getvalue())


def _seed_variety(seed_client: TestClient) -> int:
    body = seed_client.post("/api/products", json={"name": "Existente"})
    assert body.status_code == 201, body.text
    return int(body.json()["id"])


def _seed_variety_with_price(seed_client: TestClient) -> None:
    variety_id = _seed_variety(seed_client)
    body = seed_client.post(
        f"/api/products/{variety_id}/prices",
        json={"effective_from": "2026-01-01", "list_price": 1000},
    )
    assert body.status_code == 201, body.text


def _seed_stock_entry(seed_client: TestClient) -> None:
    variety_id = _seed_variety(seed_client)
    body = seed_client.post(
        "/stock/entries",
        json={"product_id": variety_id, "origin": "purchase", "quantity": 5, "total_cost": 500},
    )
    assert body.status_code == 201, body.text


def _seed_customer(seed_client: TestClient) -> None:
    body = seed_client.post(
        "/customers", json={"name": "Existente", "phone": "+5491199999999", "vip": False}
    )
    assert body.status_code == 201, body.text


def _seed_sale(seed_client: TestClient) -> None:
    variety_id = _seed_variety(seed_client)
    body = seed_client.post(
        "/sales",
        json={
            "date": "2026-01-01",
            "lines": [{"product_id": variety_id, "quantity": 1, "gift": True}],
        },
    )
    assert body.status_code == 201, body.text


def _seed_payment(seed_client: TestClient) -> None:
    customer_body = seed_client.post(
        "/customers", json={"name": "Existente", "phone": "+5491199999998", "vip": False}
    )
    customer_id = customer_body.json()["id"]
    variety_id = _seed_variety(seed_client)
    sale_body = seed_client.post(
        "/sales",
        json={
            "date": "2026-01-01",
            "customer_id": customer_id,
            "lines": [{"product_id": variety_id, "quantity": 1, "special_price": 1000}],
            "payment": "owing",
        },
    )
    assert sale_body.status_code == 201, sale_body.text
    payment_body = seed_client.post(
        f"/customers/{customer_id}/payments",
        json={"paid_on": "2026-01-01", "amount": 1000, "method": "cash"},
    )
    assert payment_body.status_code == 201, payment_body.text


def _seed_expense(seed_client: TestClient) -> None:
    category_body = seed_client.post("/expense-categories", json={"name": "Varios"})
    category_id = category_body.json()["id"]
    body = seed_client.post(
        "/expenses",
        json={
            "spent_on": "2026-01-01",
            "category_id": category_id,
            "amount": 100,
            "detail": "algo",
        },
    )
    assert body.status_code == 201, body.text


# Price entry is not independently seedable: a price entry can never outlive its
# variety (a foreign key), so varieties.has_records() already covers both at once;
# `_seed_variety_with_price` seeds them together rather than pretending they split.
_ELIGIBILITY_SEEDS = (
    _seed_variety,
    _seed_variety_with_price,
    _seed_stock_entry,
    _seed_customer,
    _seed_sale,
    _seed_payment,
    _seed_expense,
)


def test_eligibility_refuses_each_of_the_seven_kinds(
    app: FastAPI,
    engine: Engine,
    clock: MovableClock,
    tmp_path: Path,
    operator_text_en: OperatorText,
) -> None:
    for index, seed in enumerate(_ELIGIBILITY_SEEDS):
        name = f"Negocio {index}"
        with TestClient(app) as seed_client:
            token = open_business(seed_client, engine, clock, name, f"u{index}@example.com")
            signed_in(seed_client, token)
            seed(seed_client)

        business_id = _business_id(engine, name=name)
        xlsx = _save_workbook(_workbook(), tmp_path)
        mapping = _save_mapping(tmp_path, ((CUSTOMER, PHONE),), name=f"mapping-{index}.csv")
        out, err = io.StringIO(), io.StringIO()

        code = run(
            [str(business_id), str(xlsx), "--mapping", str(mapping)],
            engine=engine,
            clock=clock,
            out=out,
            err=err,
            text=operator_text_en,
        )

        assert code == 1, f"{seed.__name__}: {out.getvalue()}"
        assert "import.business_has_records" in out.getvalue()


def test_a_business_with_only_renamed_or_removed_seeded_categories_is_accepted(
    client: TestClient,
    engine: Engine,
    clock: MovableClock,
    tmp_path: Path,
    operator_text_en: OperatorText,
) -> None:
    categories = client.get("/expense-categories").json()
    by_name = {c["name"]: c["id"] for c in categories}
    client.patch(f"/expense-categories/{by_name['Other']}", json={"name": "Varios"})
    client.delete(f"/expense-categories/{by_name['Packaging']}")

    business_id = _business_id(engine)
    xlsx = _save_workbook(_workbook(), tmp_path)
    mapping = _save_mapping(tmp_path, ((CUSTOMER, PHONE),))
    out, err = io.StringIO(), io.StringIO()

    code = run(
        [str(business_id), str(xlsx), "--mapping", str(mapping)],
        engine=engine,
        clock=clock,
        out=out,
        err=err,
        text=operator_text_en,
    )

    assert code == 0, err.getvalue()
    assert len(client.get("/api/products").json()) == 1


TIE_VARIETY = "OG Kush"
CUSTOMER2 = "Maria Lopez"
PHONE2 = "1123456788"


def test_recording_order_follows_date_then_sheet_row(
    client: TestClient,
    engine: Engine,
    clock: MovableClock,
    tmp_path: Path,
    operator_text_en: OperatorText,
) -> None:
    """Two things, both end to end through `cli.run`:

    1. Across dates, date decides over sheet order (`VARIETY`/`CUSTOMER`): Entradas and
       Ventas each list a later date above an earlier one, and a negative `Ajuste`,
       dated the same day as the earlier sale, is sized to drain enough of the cheaper
       lot that the earlier sale spans both lots — a line cost and an `owing_since`
       that would come out differently had sheet order, not date, decided what got
       recorded first.
    2. On a tied date, sheet order breaks the tie (`TIE_VARIETY`/`CUSTOMER2`, kept
       apart from (1) so neither touches the other's lots, balance or account): two
       Entradas share one date at different costs per gram, and two Ventas of one
       customer share another date. The row listed first in the sheet is recorded
       first — the cheaper-per-gram sale line and the order of the account rows below
       would both come out reversed had the sheet's own row order been ignored."""
    workbook = Workbook()
    active = workbook.active
    assert active is not None
    workbook.remove(active)

    ajustes = workbook.create_sheet(layout.AJUSTES_SHEET)
    ajustes.cell(row=7, column=1, value=VARIETY)
    ajustes.cell(row=8, column=1, value=TIE_VARIETY)
    ajustes.cell(row=37, column=1, value=VARIETY)
    ajustes.cell(row=37, column=2, value=date(2026, 1, 1))
    ajustes.cell(row=37, column=3, value=1000)
    ajustes.cell(row=38, column=1, value=TIE_VARIETY)
    ajustes.cell(row=38, column=2, value=date(2026, 1, 1))
    ajustes.cell(row=38, column=3, value=1000)

    workbook.create_sheet(layout.CATALOGO_SHEET)

    entradas = workbook.create_sheet(layout.ENTRADAS_SHEET)
    _titles(entradas, ENTRADAS_TITLES)
    # -- (1) Across dates: date decides. -------------------------------------------
    # Row 2: the LATER-dated, costlier lot, listed FIRST in the sheet.
    entradas.cell(row=2, column=1, value=date(2026, 1, 10))
    entradas.cell(row=2, column=2, value=VARIETY)
    entradas.cell(row=2, column=3, value=10)
    entradas.cell(row=2, column=4, value="Compra")
    entradas.cell(row=2, column=5, value=1000)  # 100/g
    # Row 3: the EARLIER-dated, cheaper lot, listed SECOND — date must still win.
    entradas.cell(row=3, column=1, value=date(2026, 1, 1))
    entradas.cell(row=3, column=2, value=VARIETY)
    entradas.cell(row=3, column=3, value=10)
    entradas.cell(row=3, column=4, value="Compra")
    entradas.cell(row=3, column=5, value=500)  # 50/g
    # Row 4: a negative adjustment on the earlier sale's own date, draining 8 g of the
    # cheap lot. Entradas write before any Ventas, so this always gets a `recorded_at`
    # earlier than that sale's, whatever rank alone would otherwise break a tie on.
    entradas.cell(row=4, column=1, value=date(2026, 1, 15))
    entradas.cell(row=4, column=2, value=VARIETY)
    entradas.cell(row=4, column=3, value=-8)
    entradas.cell(row=4, column=4, value="Ajuste")
    # -- (2) On a tied date: sheet order decides. ------------------------------------
    # Row 5: listed FIRST, so it is recorded first and becomes the older lot, despite
    # being the costlier one per gram.
    entradas.cell(row=5, column=1, value=date(2026, 3, 1))
    entradas.cell(row=5, column=2, value=TIE_VARIETY)
    entradas.cell(row=5, column=3, value=10)
    entradas.cell(row=5, column=4, value="Compra")
    entradas.cell(row=5, column=5, value=1000)  # 100/g
    # Row 6: same date, listed SECOND — the newer lot, whatever its own cost.
    entradas.cell(row=6, column=1, value=date(2026, 3, 1))
    entradas.cell(row=6, column=2, value=TIE_VARIETY)
    entradas.cell(row=6, column=3, value=10)
    entradas.cell(row=6, column=4, value="Compra")
    entradas.cell(row=6, column=5, value=500)  # 50/g

    gastos = workbook.create_sheet(layout.GASTOS_SHEET)
    _titles(gastos, GASTOS_TITLES)

    ventas = workbook.create_sheet(layout.VENTAS_SHEET)
    _titles(ventas, VENTAS_TITLES)
    # -- (1) Across dates: date decides. -------------------------------------------
    # Row 2: the LATER-dated sale, listed FIRST in the sheet.
    ventas.cell(row=2, column=1, value=date(2026, 1, 20))
    ventas.cell(row=2, column=2, value=VARIETY)
    ventas.cell(row=2, column=3, value=5)
    ventas.cell(row=2, column=4, value="Pendiente")
    ventas.cell(row=2, column=5, value=5000)
    ventas.cell(row=2, column=6, value=CUSTOMER)
    ventas.cell(row=2, column=7, value="")
    # Row 3: the EARLIER-dated sale, listed SECOND — date must still win. It shares
    # the adjustment's own date, so the adjustment lands before it.
    ventas.cell(row=3, column=1, value=date(2026, 1, 15))
    ventas.cell(row=3, column=2, value=VARIETY)
    ventas.cell(row=3, column=3, value=3)
    ventas.cell(row=3, column=4, value="Pendiente")
    ventas.cell(row=3, column=5, value=3000)
    ventas.cell(row=3, column=6, value=CUSTOMER)
    ventas.cell(row=3, column=7, value="")
    # -- (2) On a tied date: sheet order decides, for one customer's two sales. ------
    # Row 4: listed FIRST, so it draws stock and settles first.
    ventas.cell(row=4, column=1, value=date(2026, 3, 10))
    ventas.cell(row=4, column=2, value=TIE_VARIETY)
    ventas.cell(row=4, column=3, value=5)
    ventas.cell(row=4, column=4, value="Pendiente")
    ventas.cell(row=4, column=5, value=5000)
    ventas.cell(row=4, column=6, value=CUSTOMER2)
    ventas.cell(row=4, column=7, value="")
    # Row 5: same date, listed SECOND.
    ventas.cell(row=5, column=1, value=date(2026, 3, 10))
    ventas.cell(row=5, column=2, value=TIE_VARIETY)
    ventas.cell(row=5, column=3, value=3)
    ventas.cell(row=5, column=4, value="Pendiente")
    ventas.cell(row=5, column=5, value=3000)
    ventas.cell(row=5, column=6, value=CUSTOMER2)
    ventas.cell(row=5, column=7, value="")

    xlsx = _save_workbook(workbook, tmp_path, name="order.xlsx")
    mapping = _save_mapping(
        tmp_path, ((CUSTOMER, PHONE), (CUSTOMER2, PHONE2)), name="order-mapping.csv"
    )
    business_id = _business_id(engine)
    out, err = io.StringIO(), io.StringIO()

    code = run(
        [str(business_id), str(xlsx), "--mapping", str(mapping)],
        engine=engine,
        clock=clock,
        out=out,
        err=err,
        text=operator_text_en,
    )
    assert code == 0, err.getvalue()

    stock = {row["product_id"]: row for row in client.get("/stock").json()["products"]}
    varieties_by_name = {v["name"]: v["id"] for v in client.get("/api/products").json()}
    # (1) 20 g bought in, 8 g adjusted out, 3 g + 5 g sold: 4 g left, all in the 100/g
    # lot.
    assert stock[varieties_by_name[VARIETY]]["quantity"] == 4.0
    # (2) 20 g bought in, 5 g + 3 g sold: 12 g left.
    assert stock[varieties_by_name[TIE_VARIETY]]["quantity"] == 12.0

    customers_by_name = {c["name"]: c for c in client.get("/customers").json()}
    assert customers_by_name[CUSTOMER]["balance"] == 8000
    # The earliest unpaid sale is the 01-15 one (row 3), not row 2's 01-20, whatever
    # the sheet's own order was.
    assert customers_by_name[CUSTOMER]["owing_since"] == "2026-01-15"
    assert customers_by_name[CUSTOMER2]["balance"] == 8000
    assert customers_by_name[CUSTOMER2]["owing_since"] == "2026-03-10"

    account = client.get(f"/customers/{customers_by_name[CUSTOMER]['id']}/account").json()
    sale_rows = [row for row in account["rows"] if row["kind"] == "sale"]
    # The earlier-dated sale (3000, row 3) settles first, even though row 2 (5000)
    # comes first in the sheet.
    assert [row["sale"]["total"] for row in sale_rows] == [3000, 5000]
    assert [row["running_balance"] for row in sale_rows] == [3000, 8000]

    sale_ids = [row["sale"]["id"] for row in sale_rows]
    earliest_detail = client.get(f"/sales/{sale_ids[0]}").json()
    latest_detail = client.get(f"/sales/{sale_ids[1]}").json()
    # The adjustment drains 8 g of the 10 g, 50/g lot first (it is recorded before
    # every sale on this variety, and the Entradas-before-Ventas write order already
    # gives it an earlier `recorded_at` than the earliest sale's own), leaving it only
    # 2 g: the earliest sale's 3 g spans both lots (2 g @ 50 + 1 g @ 100 = 200), and
    # the latest sale's 5 g comes entirely from the 100/g lot (500). Had sheet order,
    # rather than date, decided which lot came first, or had the adjustment landed
    # anywhere else, both figures would come out differently.
    assert earliest_detail["lines"][0]["cost"]["costed"] == 200
    assert latest_detail["lines"][0]["cost"]["costed"] == 500

    # (2) The tied-date pair: row 5 (first in the sheet) is the older, 100/g lot, so
    # both of customer 2's sales (5 g then 3 g, 8 g total, well under its own 10 g)
    # draw from it entirely. Had row 6 (50/g) been recorded first instead, both lines
    # below would cost half as much.
    tie_account = client.get(f"/customers/{customers_by_name[CUSTOMER2]['id']}/account").json()
    tie_sale_rows = [row for row in tie_account["rows"] if row["kind"] == "sale"]
    # Row 4 (5000) is recorded and settles before row 5 (3000), in sheet order.
    assert [row["sale"]["total"] for row in tie_sale_rows] == [5000, 3000]
    assert [row["running_balance"] for row in tie_sale_rows] == [5000, 8000]
    tie_sale_ids = [row["sale"]["id"] for row in tie_sale_rows]
    first_tie_detail = client.get(f"/sales/{tie_sale_ids[0]}").json()
    second_tie_detail = client.get(f"/sales/{tie_sale_ids[1]}").json()
    assert first_tie_detail["lines"][0]["cost"]["costed"] == 500
    assert second_tie_detail["lines"][0]["cost"]["costed"] == 300


def test_main_takes_its_business_day_from_the_clock_it_is_given(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    clock: MovableClock,
) -> None:
    url = f"sqlite:///{tmp_path / 'apostrophe.db'}"
    monkeypatch.setenv("APOSTROPHE_SECRET_KEY", "cli-secret")
    monkeypatch.setenv("APOSTROPHE_DATABASE_URL", url)
    file_engine = make_engine(url)
    file_app = create_app(
        Settings("cli-secret", url),
        engine=file_engine,
        hasher=TEST_HASHER,
        telegram=FakeTelegram(),
        upkeep_ticker=False,
        clock=clock,
    )
    with TestClient(file_app) as seed_client:
        open_business(seed_client, file_engine, clock, "Primer Negocio", "uno@example.com")
    business_id = _business_id(file_engine)
    xlsx = _save_workbook(_workbook(), tmp_path)
    mapping = _save_mapping(tmp_path, ((CUSTOMER, PHONE),))
    argv = [str(business_id), str(xlsx), "--mapping", str(mapping), "--dry-run"]

    # The sheet's rows are dated in January 2026: past on the app's clock, still to come
    # on a clock set in the December before.
    assert main(argv, clock=clock) == 0
    assert "Writable: yes" in capsys.readouterr().out
    assert main(argv, clock=MovableClock(datetime(2025, 12, 31, 12, 0, tzinfo=UTC))) == 0
    assert "Writable: no" in capsys.readouterr().out


def test_an_invalid_forced_language_prints_only_refused_invalid_setting(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    clock: MovableClock,
) -> None:
    monkeypatch.setenv("APOSTROPHE_SECRET_KEY", "cli-secret")
    monkeypatch.setenv("APOSTROPHE_DATABASE_URL", f"sqlite:///{tmp_path / 'apostrophe.db'}")
    monkeypatch.setenv("APOSTROPHE_FORCED_LANGUAGE", "de")

    assert main(["1", str(tmp_path / "sheet.xlsx")], clock=clock) == 1

    captured = capsys.readouterr()
    assert captured.err == "refused: InvalidSetting\n"
    assert captured.out == ""
    assert not (tmp_path / "apostrophe.db").exists()


def test_a_refusal_the_translator_cannot_render_prints_the_fixed_english_line(
    operator_text_en: OperatorText,
) -> None:
    broken = Settings("k", locales_dir=Path(__file__).resolve().parents[1] / "fixtures" / "locales")
    out, err = io.StringIO(), io.StringIO()
    engine = make_engine("sqlite://")

    code = run(
        ["1", "missing.xlsx"],
        engine=engine,
        clock=MovableClock(datetime(2026, 10, 2, 18, 0, tzinfo=UTC)),
        out=out,
        err=err,
        text=broken.operator_translator(),
    )

    assert code == 1
    assert err.getvalue().startswith("refused: ")
    assert "\n" not in err.getvalue().rstrip("\n")
