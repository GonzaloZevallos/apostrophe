"""`read_workbook` (CODING_STANDARDS.md section 5: tests at the seams, no
infrastructure beyond the `.xlsx` file itself). Every fixture is built here with
openpyxl; the real sheet and real customer data never enter the repository.
"""

from __future__ import annotations

from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.shared.grams import Grams
from app.spreadsheet_import.adapters import layout
from app.spreadsheet_import.adapters.cells import read_workbook
from app.spreadsheet_import.domain.rows import ParsedWorkbook

type _CellValue = int | float | str | date | datetime | time | None

VENTAS_TITLES = (
    "Fecha",
    "Variedad",
    "Gramos",
    "Pago",
    "Total",
    "Cliente",
    "Tipo",
    "Envío cobrado",
    "Precio especial (total)",
    "Nota",
)
ENTRADAS_TITLES = ("Fecha", "Variedad", "Gramos", "Origen", "Costo total", "Nota", "Revisar")
GASTOS_TITLES = ("Fecha", "Categoría", "Monto", "Detalle", "Revisar")


def _blank_workbook() -> Workbook:
    workbook = Workbook()
    active = workbook.active
    assert active is not None
    workbook.remove(active)
    workbook.create_sheet(layout.AJUSTES_SHEET)
    workbook.create_sheet(layout.CATALOGO_SHEET)
    workbook.create_sheet(layout.VENTAS_SHEET)
    workbook.create_sheet(layout.ENTRADAS_SHEET)
    workbook.create_sheet(layout.GASTOS_SHEET)
    return workbook


def _titles(ws: Worksheet, titles: tuple[str, ...]) -> None:
    for index, title in enumerate(titles, start=1):
        ws.cell(row=layout.TITLE_ROW, column=index, value=title)


def _row(ws: Worksheet, row: int, values: tuple[_CellValue, ...]) -> None:
    for index, value in enumerate(values, start=1):
        ws.cell(row=row, column=index, value=value)


def _save(workbook: Workbook, tmp_path: Path) -> Path:
    path = tmp_path / "sheet.xlsx"
    workbook.save(path)
    return path


def _default(tmp_path: Path) -> tuple[Workbook, Path]:
    workbook = _blank_workbook()
    _titles(workbook[layout.VENTAS_SHEET], VENTAS_TITLES)
    _titles(workbook[layout.ENTRADAS_SHEET], ENTRADAS_TITLES)
    _titles(workbook[layout.GASTOS_SHEET], GASTOS_TITLES)
    return workbook, _save(workbook, tmp_path)


def test_float_grams_with_two_decimals_are_read_as_grams(tmp_path: Path) -> None:
    workbook = _blank_workbook()
    _titles(workbook[layout.VENTAS_SHEET], VENTAS_TITLES)
    _titles(workbook[layout.GASTOS_SHEET], GASTOS_TITLES)
    entradas = workbook[layout.ENTRADAS_SHEET]
    _titles(entradas, ENTRADAS_TITLES)
    _row(entradas, 2, (date(2026, 1, 5), "Jack Herer", 3.3, "Compra", 1000, "", ""))
    _row(entradas, 3, (date(2026, 1, 6), "Jack Herer", 0.1, "Compra", 50, "", ""))

    parsed = read_workbook(_save(workbook, tmp_path))

    assert [e.grams for e in parsed.entradas] == [Grams.of("3.3"), Grams.of("0.1")]
    assert parsed.problems == ()


def test_grams_with_more_than_two_decimals_is_a_problem(tmp_path: Path) -> None:
    workbook = _blank_workbook()
    _titles(workbook[layout.VENTAS_SHEET], VENTAS_TITLES)
    _titles(workbook[layout.GASTOS_SHEET], GASTOS_TITLES)
    entradas = workbook[layout.ENTRADAS_SHEET]
    _titles(entradas, ENTRADAS_TITLES)
    _row(entradas, 2, (date(2026, 1, 5), "Jack Herer", 3.333, "Compra", 1000, "", ""))

    parsed = read_workbook(_save(workbook, tmp_path))

    assert parsed.entradas[0].grams is None
    assert len(parsed.problems) == 1
    problem = parsed.problems[0]
    assert problem.column == "Gramos"
    assert problem.kind == "grams_too_precise"
    assert problem.ref.sheet == layout.ENTRADAS_SHEET
    assert problem.ref.row == 2


def test_whole_and_fractional_pesos(tmp_path: Path) -> None:
    workbook = _blank_workbook()
    _titles(workbook[layout.VENTAS_SHEET], VENTAS_TITLES)
    gastos = workbook[layout.GASTOS_SHEET]
    _titles(gastos, GASTOS_TITLES)
    _row(gastos, 2, (date(2026, 1, 5), "Insumos", 100.0, "Bolsas", ""))
    _row(gastos, 3, (date(2026, 1, 6), "Insumos", 100.5, "Bolsas", ""))
    _titles(workbook[layout.ENTRADAS_SHEET], ENTRADAS_TITLES)

    parsed = read_workbook(_save(workbook, tmp_path))

    assert parsed.gastos[0].amount == 100
    assert parsed.gastos[1].amount is None
    assert len(parsed.problems) == 1
    problem = parsed.problems[0]
    assert problem.column == "Monto"
    assert problem.kind == "pesos_not_integer"
    assert problem.ref.row == 3


def test_dates_and_datetimes_at_midnight_are_both_read(tmp_path: Path) -> None:
    workbook = _blank_workbook()
    _titles(workbook[layout.VENTAS_SHEET], VENTAS_TITLES)
    _titles(workbook[layout.GASTOS_SHEET], GASTOS_TITLES)
    entradas = workbook[layout.ENTRADAS_SHEET]
    _titles(entradas, ENTRADAS_TITLES)
    _row(entradas, 2, (date(2026, 1, 5), "Jack Herer", 1, "Compra", 10, "", ""))
    _row(entradas, 3, (datetime(2026, 1, 6, 0, 0, 0), "Jack Herer", 1, "Compra", 10, "", ""))

    parsed = read_workbook(_save(workbook, tmp_path))

    assert [e.entry_date for e in parsed.entradas] == [date(2026, 1, 5), date(2026, 1, 6)]
    assert parsed.problems == ()


def test_a_datetime_off_midnight_is_a_problem(tmp_path: Path) -> None:
    workbook = _blank_workbook()
    _titles(workbook[layout.VENTAS_SHEET], VENTAS_TITLES)
    _titles(workbook[layout.GASTOS_SHEET], GASTOS_TITLES)
    entradas = workbook[layout.ENTRADAS_SHEET]
    _titles(entradas, ENTRADAS_TITLES)
    _row(entradas, 2, (datetime(2026, 1, 5, 10, 30), "Jack Herer", 1, "Compra", 10, "", ""))

    parsed = read_workbook(_save(workbook, tmp_path))

    assert parsed.entradas[0].entry_date is None
    assert len(parsed.problems) == 1
    assert parsed.problems[0].kind == "date_invalid"
    assert parsed.problems[0].column == "Fecha"


def test_each_missing_title_is_reported_and_the_sheet_yields_no_rows(tmp_path: Path) -> None:
    workbook = _blank_workbook()
    ventas = workbook[layout.VENTAS_SHEET]
    _titles(ventas, tuple(t for t in VENTAS_TITLES if t not in ("Cliente", "Nota")))
    _row(ventas, 2, (date(2026, 1, 5), "Jack Herer", 1, "Efectivo", 1000, "Lista", 0, ""))
    _titles(workbook[layout.ENTRADAS_SHEET], ENTRADAS_TITLES)
    _titles(workbook[layout.GASTOS_SHEET], GASTOS_TITLES)

    parsed = read_workbook(_save(workbook, tmp_path))

    missing = {p.column for p in parsed.problems if p.kind == "column_missing"}
    assert missing == {"Cliente", "Nota"}
    assert parsed.ventas == ()


def test_a_duplicated_title_is_reported_and_the_sheet_yields_no_rows(tmp_path: Path) -> None:
    workbook = _blank_workbook()
    ventas = workbook[layout.VENTAS_SHEET]
    _titles(ventas, VENTAS_TITLES + ("Fecha",))
    _titles(workbook[layout.ENTRADAS_SHEET], ENTRADAS_TITLES)
    _titles(workbook[layout.GASTOS_SHEET], GASTOS_TITLES)

    parsed = read_workbook(_save(workbook, tmp_path))

    duplicated = [p for p in parsed.problems if p.kind == "column_duplicate"]
    assert [p.column for p in duplicated] == ["Fecha"]
    assert parsed.ventas == ()


def test_a_missing_entradas_title_is_reported_and_the_sheet_yields_no_rows(tmp_path: Path) -> None:
    workbook = _blank_workbook()
    _titles(workbook[layout.VENTAS_SHEET], VENTAS_TITLES)
    _titles(workbook[layout.GASTOS_SHEET], GASTOS_TITLES)
    entradas = workbook[layout.ENTRADAS_SHEET]
    _titles(entradas, tuple(t for t in ENTRADAS_TITLES if t not in ("Origen", "Revisar")))
    _row(entradas, 2, (date(2026, 1, 5), "Jack Herer", 1, 1000, ""))

    parsed = read_workbook(_save(workbook, tmp_path))

    missing = {p.column for p in parsed.problems if p.kind == "column_missing"}
    assert missing == {"Origen", "Revisar"}
    assert parsed.entradas == ()


def test_a_duplicated_gastos_title_is_reported_and_the_sheet_yields_no_rows(tmp_path: Path) -> None:
    workbook = _blank_workbook()
    _titles(workbook[layout.VENTAS_SHEET], VENTAS_TITLES)
    _titles(workbook[layout.ENTRADAS_SHEET], ENTRADAS_TITLES)
    gastos = workbook[layout.GASTOS_SHEET]
    _titles(gastos, GASTOS_TITLES + ("Monto",))

    parsed = read_workbook(_save(workbook, tmp_path))

    duplicated = [p for p in parsed.problems if p.kind == "column_duplicate"]
    assert [p.column for p in duplicated] == ["Monto"]
    assert parsed.gastos == ()


def test_trailing_formula_only_rows_do_not_appear(tmp_path: Path) -> None:
    workbook = _blank_workbook()
    ventas = workbook[layout.VENTAS_SHEET]
    _titles(ventas, VENTAS_TITLES)
    _row(
        ventas, 2, (date(2026, 1, 5), "Jack Herer", 1, "Efectivo", 1000, "Ana", "Lista", 0, "", "")
    )
    total_column = VENTAS_TITLES.index("Total") + 1
    for offset in range(50):
        # Typed columns blank; only the formula column (Total) carries a cached
        # value, the way a spreadsheet's filled-down formula would.
        ventas.cell(row=3 + offset, column=total_column, value="Falta fecha")
    _titles(workbook[layout.ENTRADAS_SHEET], ENTRADAS_TITLES)
    _titles(workbook[layout.GASTOS_SHEET], GASTOS_TITLES)

    parsed = read_workbook(_save(workbook, tmp_path))

    assert len(parsed.ventas) == 1
    assert parsed.problems == ()


def test_a_ventas_total_with_no_cached_value_blocks_once(tmp_path: Path) -> None:
    workbook = _blank_workbook()
    ventas = workbook[layout.VENTAS_SHEET]
    _titles(ventas, VENTAS_TITLES)
    total_column = VENTAS_TITLES.index("Total") + 1
    _row(
        ventas, 2, (date(2026, 1, 5), "Jack Herer", 1, "Efectivo", None, "Ana", "Lista", 0, "", "")
    )
    _row(
        ventas, 3, (date(2026, 1, 6), "Jack Herer", 1, "Efectivo", None, "Ana", "Lista", 0, "", "")
    )
    # A real formula, exactly as the sheet's own Total column holds one. openpyxl
    # never evaluates it, so re-reading it with `data_only=True` gives `None`: a
    # formula with no cached value, not a cell deliberately left blank.
    for row in (2, 3):
        ventas.cell(row=row, column=total_column, value="=C2*2")
    _titles(workbook[layout.ENTRADAS_SHEET], ENTRADAS_TITLES)
    _titles(workbook[layout.GASTOS_SHEET], GASTOS_TITLES)

    parsed = read_workbook(_save(workbook, tmp_path))

    assert len(parsed.ventas) == 2
    assert [v.total for v in parsed.ventas] == [None, None]
    not_recalculated = [p for p in parsed.problems if p.kind == "not_recalculated"]
    assert len(not_recalculated) == 1


def test_a_blank_row_is_skipped_before_its_formula_cell_is_read(tmp_path: Path) -> None:
    workbook = _blank_workbook()
    ventas = workbook[layout.VENTAS_SHEET]
    _titles(ventas, VENTAS_TITLES)
    total_column = VENTAS_TITLES.index("Total") + 1
    # Typed columns wholly blank; the Total formula still shows its "no date" text,
    # as it would over a genuinely empty sheet row.
    ventas.cell(row=2, column=total_column, value="Falta fecha")
    _titles(workbook[layout.ENTRADAS_SHEET], ENTRADAS_TITLES)
    _titles(workbook[layout.GASTOS_SHEET], GASTOS_TITLES)

    parsed = read_workbook(_save(workbook, tmp_path))

    assert parsed.ventas == ()
    assert parsed.problems == ()


def test_variety_rows_and_thresholds_blank_rows_skipped(tmp_path: Path) -> None:
    workbook, _ = _default(tmp_path)
    ajustes = workbook[layout.AJUSTES_SHEET]
    ajustes.cell(row=7, column=layout.VARIETY_NAME_COLUMN, value="Jack Herer")
    ajustes.cell(row=7, column=layout.VARIETY_THRESHOLD_COLUMN, value=50)
    ajustes.cell(row=8, column=layout.VARIETY_NAME_COLUMN, value="Blue Dream")
    # Row 9 left wholly blank: skipped, no problem.

    parsed = read_workbook(_save(workbook, tmp_path))

    assert [(v.raw_name, v.threshold) for v in parsed.varieties] == [
        ("Jack Herer", Grams.of(50)),
        ("Blue Dream", None),
    ]
    assert parsed.problems == ()


def test_price_history_rows_from_row_37(tmp_path: Path) -> None:
    workbook, _ = _default(tmp_path)
    ajustes = workbook[layout.AJUSTES_SHEET]
    _row(ajustes, 37, ("Jack Herer", date(2026, 1, 1), 1000, 800, 10))

    parsed = read_workbook(_save(workbook, tmp_path))

    [entry] = parsed.price_history
    assert entry.variety == "Jack Herer"
    assert entry.effective == date(2026, 1, 1)
    assert entry.list_price == 1000
    assert entry.bulk_price == 800
    assert entry.bulk_minimum == Grams.of(10)
    assert parsed.problems == ()


def test_catalog_rows_and_header_footer(tmp_path: Path) -> None:
    workbook, _ = _default(tmp_path)
    catalogo = workbook[layout.CATALOGO_SHEET]
    _row(catalogo, 6, ("Jack Herer", "Sí", "🌿", 60, "Tierra"))
    catalogo.cell(row=18, column=2, value="Bienvenidos")
    catalogo.cell(row=19, column=2, value="Gracias")

    parsed = read_workbook(_save(workbook, tmp_path))

    [row] = parsed.catalog
    assert row.variety == "Jack Herer"
    assert row.shown == "Sí"
    assert row.emoji == "🌿"
    assert row.indica == 60
    assert row.aroma == "Tierra"
    assert parsed.header_footer.header == "Bienvenidos"
    assert parsed.header_footer.footer == "Gracias"


def test_header_footer_blank_cells_give_blank_text(tmp_path: Path) -> None:
    workbook, _ = _default(tmp_path)

    parsed = read_workbook(_save(workbook, tmp_path))

    assert parsed.header_footer.header == ""
    assert parsed.header_footer.footer == ""


def test_categories_rows_a21_a28(tmp_path: Path) -> None:
    workbook, _ = _default(tmp_path)
    ajustes = workbook[layout.AJUSTES_SHEET]
    ajustes.cell(row=21, column=1, value="Insumos")
    ajustes.cell(row=22, column=1, value="Transporte")

    parsed = read_workbook(_save(workbook, tmp_path))

    assert [c.raw_name for c in parsed.categories] == ["Insumos", "Transporte"]


def test_entradas_and_gastos_revisar_is_read_as_text_and_never_stored_by_shape(
    tmp_path: Path,
) -> None:
    workbook = _blank_workbook()
    _titles(workbook[layout.VENTAS_SHEET], VENTAS_TITLES)
    entradas = workbook[layout.ENTRADAS_SHEET]
    _titles(entradas, ENTRADAS_TITLES)
    _row(entradas, 2, (date(2026, 1, 5), "Jack Herer", 1, "Compra", 10, "", "Falta costo"))
    _row(entradas, 3, (date(2026, 1, 6), "Jack Herer", 1, "Compra", 10, "", ""))
    gastos = workbook[layout.GASTOS_SHEET]
    _titles(gastos, GASTOS_TITLES)
    _row(gastos, 2, (date(2026, 1, 5), "Insumos", 10, "Bolsas", "Falta categoría"))

    parsed = read_workbook(_save(workbook, tmp_path))

    assert [e.revisar for e in parsed.entradas] == ["Falta costo", None]
    assert [g.revisar for g in parsed.gastos] == ["Falta categoría"]
    assert parsed.problems == ()


def test_text_columns_hold_trimmed_text(tmp_path: Path) -> None:
    workbook = _blank_workbook()
    _titles(workbook[layout.ENTRADAS_SHEET], ENTRADAS_TITLES)
    _titles(workbook[layout.GASTOS_SHEET], GASTOS_TITLES)
    ventas = workbook[layout.VENTAS_SHEET]
    _titles(ventas, VENTAS_TITLES)
    _row(
        ventas,
        2,
        (
            date(2026, 1, 5),
            "  Jack Herer  ",
            1,
            "Efectivo",
            1000,
            "Ana",
            "Lista",
            0,
            "",
            "  nota  ",
        ),
    )

    parsed = read_workbook(_save(workbook, tmp_path))

    [row] = parsed.ventas
    assert row.variety == "Jack Herer"
    assert row.note == "nota"


def test_read_workbook_returns_a_parsed_workbook(tmp_path: Path) -> None:
    workbook, _ = _default(tmp_path)

    parsed = read_workbook(_save(workbook, tmp_path))

    assert isinstance(parsed, ParsedWorkbook)


# -- GH-55: whole-number floats, numbers and dates in text cells ----------------------


def test_a_whole_float_share_and_whole_float_pesos_are_accepted(tmp_path: Path) -> None:
    workbook, _ = _default(tmp_path)
    catalogo = workbook[layout.CATALOGO_SHEET]
    _row(catalogo, 6, ("Jack Herer", "Sí", "🌿", 60.0, "Tierra"))
    gastos = workbook[layout.GASTOS_SHEET]
    _row(gastos, 2, (date(2026, 1, 5), "Insumos", 500.0, "Bolsas", ""))

    parsed = read_workbook(_save(workbook, tmp_path))

    [catalog_row] = parsed.catalog
    assert catalog_row.indica == 60
    assert isinstance(catalog_row.indica, int)
    assert parsed.gastos[0].amount == 500
    assert isinstance(parsed.gastos[0].amount, int)
    assert parsed.problems == ()


def test_a_non_whole_float_share_is_refused_by_the_catalog_owner_not_the_reader(
    tmp_path: Path,
) -> None:
    workbook, _ = _default(tmp_path)
    catalogo = workbook[layout.CATALOGO_SHEET]
    _row(catalogo, 6, ("Jack Herer", "Sí", "🌿", 60.5, "Tierra"))

    parsed = read_workbook(_save(workbook, tmp_path))

    [catalog_row] = parsed.catalog
    assert catalog_row.indica == "60.5"
    assert parsed.problems == ()


def test_a_ventas_total_non_whole_float_is_kept_as_an_exact_decimal(tmp_path: Path) -> None:
    workbook = _blank_workbook()
    _titles(workbook[layout.ENTRADAS_SHEET], ENTRADAS_TITLES)
    _titles(workbook[layout.GASTOS_SHEET], GASTOS_TITLES)
    ventas = workbook[layout.VENTAS_SHEET]
    _titles(ventas, VENTAS_TITLES)
    _row(
        ventas,
        2,
        (date(2026, 1, 5), "Jack Herer", 1, "Efectivo", 15501.55, "Ana", "Lista", 0, "", ""),
    )

    parsed = read_workbook(_save(workbook, tmp_path))

    [row] = parsed.ventas
    assert row.total == Decimal("15501.55")
    assert parsed.problems == ()


def test_a_number_in_a_text_column_reads_back_as_its_own_text(tmp_path: Path) -> None:
    workbook = _blank_workbook()
    _titles(workbook[layout.GASTOS_SHEET], GASTOS_TITLES)
    _titles(workbook[layout.VENTAS_SHEET], VENTAS_TITLES)
    entradas = workbook[layout.ENTRADAS_SHEET]
    _titles(entradas, ENTRADAS_TITLES)
    _row(entradas, 2, (date(2026, 1, 5), 420, 1, "Compra", 10, 100.0, ""))
    _row(entradas, 3, (date(2026, 1, 6), "Jack Herer", 1, "Compra", 10, 0.25, ""))
    _row(entradas, 4, (date(2026, 1, 7), "Jack Herer", 1, "Compra", 10, 1e16, ""))

    parsed = read_workbook(_save(workbook, tmp_path))

    assert [e.variety for e in parsed.entradas] == ["420", "Jack Herer", "Jack Herer"]
    assert [e.note for e in parsed.entradas] == ["100", "0,25", "10000000000000000"]
    assert parsed.problems == ()


def test_a_boolean_in_a_text_column_is_still_text_invalid(tmp_path: Path) -> None:
    workbook = _blank_workbook()
    _titles(workbook[layout.VENTAS_SHEET], VENTAS_TITLES)
    _titles(workbook[layout.GASTOS_SHEET], GASTOS_TITLES)
    entradas = workbook[layout.ENTRADAS_SHEET]
    _titles(entradas, ENTRADAS_TITLES)
    _row(entradas, 2, (date(2026, 1, 5), "Jack Herer", 1, "Compra", 10, True, ""))

    parsed = read_workbook(_save(workbook, tmp_path))

    assert parsed.entradas[0].note is None
    assert len(parsed.problems) == 1
    assert parsed.problems[0].kind == "text_invalid"
    assert parsed.problems[0].column == "Nota"


def test_a_date_in_a_text_column_is_text_is_date_not_text_invalid(tmp_path: Path) -> None:
    workbook = _blank_workbook()
    _titles(workbook[layout.VENTAS_SHEET], VENTAS_TITLES)
    _titles(workbook[layout.GASTOS_SHEET], GASTOS_TITLES)
    entradas = workbook[layout.ENTRADAS_SHEET]
    _titles(entradas, ENTRADAS_TITLES)
    _row(entradas, 2, (date(2026, 1, 5), "Jack Herer", 1, "Compra", 10, date(2026, 10, 12), ""))

    parsed = read_workbook(_save(workbook, tmp_path))

    assert parsed.entradas[0].note is None
    assert len(parsed.problems) == 1
    assert parsed.problems[0].kind == "text_is_date"
    assert parsed.problems[0].column == "Nota"


def test_a_time_in_a_text_column_is_also_text_is_date(tmp_path: Path) -> None:
    workbook = _blank_workbook()
    _titles(workbook[layout.VENTAS_SHEET], VENTAS_TITLES)
    _titles(workbook[layout.GASTOS_SHEET], GASTOS_TITLES)
    entradas = workbook[layout.ENTRADAS_SHEET]
    _titles(entradas, ENTRADAS_TITLES)
    _row(entradas, 2, (date(2026, 1, 5), "Jack Herer", 1, "Compra", 10, time(12, 30), ""))

    parsed = read_workbook(_save(workbook, tmp_path))

    assert parsed.entradas[0].note is None
    assert len(parsed.problems) == 1
    assert parsed.problems[0].kind == "text_is_date"


def test_a_date_in_fecha_is_unaffected_by_the_text_is_date_rule(tmp_path: Path) -> None:
    workbook = _blank_workbook()
    _titles(workbook[layout.VENTAS_SHEET], VENTAS_TITLES)
    _titles(workbook[layout.GASTOS_SHEET], GASTOS_TITLES)
    entradas = workbook[layout.ENTRADAS_SHEET]
    _titles(entradas, ENTRADAS_TITLES)
    _row(entradas, 2, (date(2026, 1, 5), "Jack Herer", 1, "Compra", 10, "", ""))

    parsed = read_workbook(_save(workbook, tmp_path))

    assert parsed.entradas[0].entry_date == date(2026, 1, 5)
    assert parsed.problems == ()
