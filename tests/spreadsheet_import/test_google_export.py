"""The Google Sheets export shape (GH-55): no worksheet keeps a real `<dimension>`
element, and every number is stored as a float. Proven at `read_workbook` (no
declared size, whole-number floats, decimal fields, numbers and dates in text cells)
and at the CLI dry run / real run (full pipeline: columns found, rows read to the
last one, fixed addresses, mismatches, owner refusals). The fixture is built by this
module with openpyxl, then rewritten in place with
`tests.spreadsheet_import.factories.as_google_export`; no real sheet enters the
repository (CODING_STANDARDS.md section 5; import.md Constraints -> Test data).
"""

from __future__ import annotations

import io
from datetime import date, time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.engine import Engine
from sqlalchemy.orm import Session

from app.accounts.adapters.tables import BusinessRow
from app.shared.grams import Grams
from app.spreadsheet_import.adapters import layout
from app.spreadsheet_import.adapters.cells import read_workbook
from app.spreadsheet_import.adapters.cli import run
from app.spreadsheet_import.domain.rows import ParsedWorkbook
from tests.conftest import MovableClock
from tests.spreadsheet_import.factories import as_google_export

VENTAS_TITLES = (
    "Fecha",
    "Variedad",
    "Gramos",
    "Pago",
    "Total",
    "Cliente",
    "Tipo",
    "Envío cobrado",
    "Precio especial (total)",
    "Nota",
)
ENTRADAS_TITLES = ("Fecha", "Variedad", "Gramos", "Origen", "Costo total", "Nota", "Revisar")
GASTOS_TITLES = ("Fecha", "Categoría", "Monto", "Detalle", "Revisar")

# A numeric name, used for the variety in every sheet and for the customer, so the
# "number in a text cell" rule is proven where it matters most: a name that still has
# to match across sheets and against the mapping (import.md "Numbers where text is
# expected").
VARIETY = "420"
CUSTOMER = "420"
PHONE = "1123456789"


def _titles(ws: Worksheet, titles: tuple[str, ...]) -> None:
    for index, title in enumerate(titles, start=1):
        ws.cell(row=layout.TITLE_ROW, column=index, value=title)


def _clean_workbook() -> Workbook:
    workbook = Workbook()
    active = workbook.active
    assert active is not None
    workbook.remove(active)

    ajustes = workbook.create_sheet(layout.AJUSTES_SHEET)
    # Sentinels just outside each fixed range/column, distinct from anything a
    # correct read would produce: an off-by-one in `_at` picks one of these up
    # instead of (or as well as) the real cell, and a test then sees it.
    ajustes.cell(row=6, column=1, value="sentinel-row6")  # just above VARIETY_FIRST_ROW
    ajustes.cell(row=17, column=1, value="sentinel-row17")  # just below VARIETY_LAST_ROW
    ajustes.cell(row=7, column=5, value="sentinel-col-e")  # just left of F7 (threshold)
    ajustes.cell(row=7, column=7, value="sentinel-col-g")  # just right of F7
    ajustes.cell(row=7, column=1, value=VARIETY)
    ajustes.cell(row=7, column=6, value=20.0)  # whole-float threshold
    ajustes.cell(row=8, column=1, value="Blue Dream")
    ajustes.cell(row=8, column=6, value=7.25)  # decimal threshold, to the hundredth
    ajustes.cell(row=20, column=1, value="sentinel-row20")  # just above CATEGORY_FIRST_ROW
    ajustes.cell(row=29, column=1, value="sentinel-row29")  # just below CATEGORY_LAST_ROW
    ajustes.cell(row=21, column=1, value=100)  # a numeric category name
    ajustes.cell(row=37, column=1, value=VARIETY)
    ajustes.cell(row=37, column=2, value=date(2026, 1, 1))
    ajustes.cell(row=37, column=3, value=1000.0)  # whole-float list price
    ajustes.cell(row=37, column=4, value=800.0)  # whole-float bulk price
    ajustes.cell(row=37, column=5, value=10)

    catalogo = workbook.create_sheet(layout.CATALOGO_SHEET)
    catalogo.cell(row=5, column=1, value="sentinel-row5")  # just above CATALOG_FIRST_ROW
    catalogo.cell(row=16, column=1, value="sentinel-row16")  # just below CATALOG_LAST_ROW
    catalogo.cell(row=6, column=1, value=VARIETY)
    catalogo.cell(row=6, column=2, value="Sí")
    catalogo.cell(row=6, column=4, value=60.0)  # whole-float indica share
    catalogo.cell(row=17, column=2, value="sentinel-row17-b18")  # just above B18
    catalogo.cell(row=18, column=1, value="sentinel-col-a18")  # just left of B18
    catalogo.cell(row=18, column=3, value="sentinel-col-c18")  # just right of B18
    catalogo.cell(row=18, column=2, value="Catálogo de la casa")
    catalogo.cell(row=19, column=1, value="sentinel-col-a19")  # just left of B19
    catalogo.cell(row=19, column=3, value="sentinel-col-c19")  # just right of B19
    catalogo.cell(row=19, column=2, value="Gracias por su compra")
    catalogo.cell(row=20, column=2, value="sentinel-row20-b19")  # just below B19

    ventas = workbook.create_sheet(layout.VENTAS_SHEET)
    _titles(ventas, VENTAS_TITLES)
    ventas.cell(row=2, column=1, value=date(2026, 1, 10))
    ventas.cell(row=2, column=2, value=VARIETY)
    ventas.cell(row=2, column=3, value=10)
    ventas.cell(row=2, column=4, value="Pendiente")
    ventas.cell(row=2, column=5, value=15501.55)  # non-whole Total: a mismatch, never blocks
    ventas.cell(row=2, column=6, value=CUSTOMER)
    ventas.cell(row=2, column=7, value="")
    ventas.cell(row=2, column=8, value=0.0)  # whole-float shipping
    ventas.cell(row=2, column=9, value=500.0)  # whole-float special total
    ventas.cell(row=2, column=10, value=100.0)  # a numeric Nota
    ventas.cell(row=3, column=1, value=date(2026, 1, 11))
    ventas.cell(row=3, column=2, value=VARIETY)
    ventas.cell(row=3, column=3, value=10)
    ventas.cell(row=3, column=4, value="Pendiente")
    # The "known risk" the plan records: floating-point noise in a cached Total.
    # Still a mismatch, still never blocks.
    ventas.cell(row=3, column=5, value=9899.999999999998)
    ventas.cell(row=3, column=6, value=CUSTOMER)
    ventas.cell(row=3, column=7, value="")

    entradas = workbook.create_sheet(layout.ENTRADAS_SHEET)
    _titles(entradas, ENTRADAS_TITLES)
    entradas.cell(row=2, column=1, value=date(2026, 1, 1))
    entradas.cell(row=2, column=2, value=VARIETY)
    entradas.cell(row=2, column=3, value=3.5)  # decimal grams, to the hundredth
    entradas.cell(row=2, column=4, value="Compra")
    entradas.cell(row=2, column=5, value=20_000.0)  # whole-float cost
    entradas.cell(row=2, column=6, value="")
    # Three more rows, read because `_last_row` keeps going past row 2, each proving
    # one more number-as-text case (import.md "Numbers where text is expected").
    for extra_row, note_value in ((3, 3.5), (4, 1e-05), (5, -2.5)):
        entradas.cell(row=extra_row, column=1, value=date(2026, 1, extra_row))
        entradas.cell(row=extra_row, column=2, value=VARIETY)
        entradas.cell(row=extra_row, column=3, value=1.0)
        entradas.cell(row=extra_row, column=4, value="Compra")
        entradas.cell(row=extra_row, column=5, value=100.0)
        entradas.cell(row=extra_row, column=6, value=note_value)

    gastos = workbook.create_sheet(layout.GASTOS_SHEET)
    _titles(gastos, GASTOS_TITLES)
    gastos.cell(row=2, column=1, value=date(2026, 1, 2))
    gastos.cell(row=2, column=2, value=100)  # matches Ajustes' numeric category
    gastos.cell(row=2, column=3, value=500.0)  # whole-float amount
    gastos.cell(row=2, column=4, value=1000.0)  # a numeric Detalle

    return workbook


def _problem_workbook() -> Workbook:
    workbook = _clean_workbook()
    catalogo = workbook[layout.CATALOGO_SHEET]
    catalogo.cell(row=6, column=4, value=60.5)  # a non-whole share: still refused

    ventas = workbook[layout.VENTAS_SHEET]
    ventas.cell(row=2, column=10, value=True)  # a boolean in a Nota
    ventas.cell(row=3, column=10, value=time(12, 30))  # a time in a Nota

    entradas = workbook[layout.ENTRADAS_SHEET]
    entradas.cell(row=2, column=4, value=3.0)  # an unknown numeric Origen
    entradas.cell(row=2, column=6, value=date(2026, 10, 12))  # a date in a Nota
    entradas.cell(row=2, column=7, value=0.0)  # a numeric Revisar: text_invalid, not unfinished

    gastos = workbook[layout.GASTOS_SHEET]
    gastos.cell(row=2, column=4, value=date(2026, 10, 13))  # a date in a Detalle

    ajustes = workbook[layout.AJUSTES_SHEET]
    # "Blue Dream" (row 8) isn't mentioned anywhere else, so turning its name into a
    # date drops only that variety, with nothing else cascading from it.
    ajustes.cell(row=8, column=1, value=date(2026, 10, 14))  # a date in a name cell

    return workbook


def _save(workbook: Workbook, tmp_path: Path, *, name: str, a1_dimension: bool) -> Path:
    path = tmp_path / name
    workbook.save(path)
    as_google_export(path, a1_dimension=a1_dimension)
    return path


def _save_mapping(tmp_path: Path, name: str = "mapping.csv") -> Path:
    path = tmp_path / name
    path.write_text(f"name,phone\n{CUSTOMER},{PHONE}\n", encoding="utf-8")
    return path


def _business_id(engine: Engine, name: str = "Primer Negocio") -> int:
    with Session(engine) as session:
        row = session.scalars(select(BusinessRow).where(BusinessRow.name == name)).one()
        assert row.id is not None
        return row.id


# -- The fixture's own shape: proven before anything about the import ----------------

# Every named numeric cell the brief asks to prove reads back as a `float`: the
# share, the list and bulk prices and the bulk minimum, both thresholds, and
# Ventas'/Entradas'/Gastos' own whole-number and decimal fields.
_FLOAT_CELLS: tuple[tuple[str, int, int], ...] = (
    (layout.CATALOGO_SHEET, 6, layout.CATALOG_INDICA_COLUMN),
    (layout.AJUSTES_SHEET, 7, layout.VARIETY_THRESHOLD_COLUMN),
    (layout.AJUSTES_SHEET, 8, layout.VARIETY_THRESHOLD_COLUMN),
    (layout.AJUSTES_SHEET, 37, layout.PRICE_HISTORY_LIST_PRICE_COLUMN),
    (layout.AJUSTES_SHEET, 37, layout.PRICE_HISTORY_BULK_PRICE_COLUMN),
    (layout.AJUSTES_SHEET, 37, layout.PRICE_HISTORY_BULK_MINIMUM_COLUMN),
    (layout.VENTAS_SHEET, 2, 3),  # Gramos
    (layout.VENTAS_SHEET, 2, 5),  # Total
    (layout.VENTAS_SHEET, 2, 8),  # Envío cobrado
    (layout.VENTAS_SHEET, 2, 9),  # Precio especial (total)
    (layout.VENTAS_SHEET, 3, 3),
    (layout.VENTAS_SHEET, 3, 5),
    (layout.ENTRADAS_SHEET, 2, 3),  # Gramos
    (layout.ENTRADAS_SHEET, 2, 5),  # Costo total
    (layout.GASTOS_SHEET, 2, 3),  # Monto
)


def _cell_value(workbook: Workbook, sheet: str, row: int, column: int) -> object:
    ws = workbook[sheet]
    # An explicit `max_col` here, never the sheet's own (possibly wrong, possibly
    # capped by the A1 variant's stale dimension) `max_column`: this helper must see
    # the real cell regardless of what the fixture's own `<dimension>` claims.
    values = next(
        ws.iter_rows(min_row=row, max_row=row, min_col=1, max_col=column, values_only=True)
    )
    return values[column - 1]


_ALL_SHEETS = (
    layout.AJUSTES_SHEET,
    layout.CATALOGO_SHEET,
    layout.VENTAS_SHEET,
    layout.ENTRADAS_SHEET,
    layout.GASTOS_SHEET,
)


def _assert_fixture_shape(path: Path, *, a1_dimension: bool) -> None:
    """The proof the plan asks for, before anything about the import is asserted:
    no usable `<dimension>` on any of the five sheets the factory rewrites, and
    every named numeric cell reads back as a `float` (GH-55's reproduction, over
    the mechanism itself, not just read from the code). Raises `AssertionError`
    on a fixture that isn't Google-shaped, the same as any other assertion."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in _ALL_SHEETS:
            assert workbook[sheet].max_column == (1 if a1_dimension else None), sheet
        for sheet, row, column in _FLOAT_CELLS:
            value = _cell_value(workbook, sheet, row, column)
            assert isinstance(value, float), (sheet, row, column, value)
    finally:
        workbook.close()


def test_the_fixture_has_no_usable_dimension_and_its_numbers_read_back_as_float(
    tmp_path: Path,
) -> None:
    path = _save(_clean_workbook(), tmp_path, name="sheet.xlsx", a1_dimension=False)

    _assert_fixture_shape(path, a1_dimension=False)


def test_the_a1_variant_also_has_no_usable_dimension(tmp_path: Path) -> None:
    path = _save(_clean_workbook(), tmp_path, name="sheet.xlsx", a1_dimension=True)

    _assert_fixture_shape(path, a1_dimension=True)


def test_the_problem_fixture_also_has_no_usable_dimension_and_its_numbers_are_float(
    tmp_path: Path,
) -> None:
    path = _save(_problem_workbook(), tmp_path, name="sheet.xlsx", a1_dimension=False)

    _assert_fixture_shape(path, a1_dimension=False)
    # The problem variant's own non-whole share and numeric Origen are still floats.
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        share = _cell_value(workbook, layout.CATALOGO_SHEET, 6, layout.CATALOG_INDICA_COLUMN)
        assert isinstance(share, float)
        origin = _cell_value(workbook, layout.ENTRADAS_SHEET, 2, 4)
        assert isinstance(origin, float)
    finally:
        workbook.close()


def test_assert_fixture_shape_rejects_a_fixture_that_is_not_google_shaped(
    tmp_path: Path,
) -> None:
    """`_assert_fixture_shape` is itself a check the fixture-building tests lean on
    (an enforcement artefact): it must fail, not pass by accident, on a workbook
    that was never put through `as_google_export` and so keeps its real dimension
    and its `int` cells — the opposite of what every other test here saves."""
    path = tmp_path / "sheet.xlsx"
    _clean_workbook().save(path)  # saved plain, deliberately skipping as_google_export

    with pytest.raises(AssertionError):
        _assert_fixture_shape(path, a1_dimension=False)


# -- `read_workbook`: columns found, rows read to the last one, floats read right -----


def _assert_clean_parse(parsed: ParsedWorkbook) -> None:
    """Every "finds everything" claim the brief makes for the clean export, checked
    once so both the no-dimension and the A1-dimension variant run exactly the same
    assertions (brief Acceptance: "over both ... finds every column and reads every
    row ... every fixed address")."""
    assert {p.column for p in parsed.problems if p.kind == "column_missing"} == set()
    assert parsed.problems == ()
    assert len(parsed.ventas) == 2
    assert len(parsed.entradas) == 4
    assert len(parsed.gastos) == 1
    [history] = parsed.price_history
    assert history.list_price == 1000
    assert isinstance(history.list_price, int)
    assert history.bulk_price == 800
    assert history.bulk_minimum == Grams.of(10)
    [variety] = [v for v in parsed.varieties if v.raw_name == VARIETY]
    assert variety.threshold == Grams.of(20)
    [blue_dream] = [v for v in parsed.varieties if v.raw_name == "Blue Dream"]
    assert blue_dream.threshold == Grams.of("7.25")
    [catalog_row] = parsed.catalog
    assert catalog_row.indica == 60
    assert isinstance(catalog_row.indica, int)
    assert parsed.header_footer.header == "Catálogo de la casa"
    assert parsed.header_footer.footer == "Gracias por su compra"
    [category] = parsed.categories
    assert category.raw_name == "100"
    [first_venta] = [v for v in parsed.ventas if v.ref.row == 2]
    assert first_venta.special_total == 500
    assert isinstance(first_venta.special_total, int)
    first_entrada = next(e for e in parsed.entradas if e.ref.row == 2)
    assert first_entrada.grams == Grams.of("3.5")
    assert first_entrada.total_cost == 20_000
    notes = {e.ref.row: e.note for e in parsed.entradas}
    assert notes[3] == "3,5"
    assert notes[4] == "0,00001"
    assert notes[5] == "-2,5"
    [gasto] = parsed.gastos
    assert gasto.category == "100"
    assert gasto.amount == 500
    assert gasto.detail == "1000"


def test_read_workbook_over_the_clean_export_finds_every_column_and_row(
    tmp_path: Path,
) -> None:
    path = _save(_clean_workbook(), tmp_path, name="sheet.xlsx", a1_dimension=False)
    _assert_fixture_shape(path, a1_dimension=False)

    parsed = read_workbook(path)

    _assert_clean_parse(parsed)


def test_read_workbook_matches_the_numeric_variety_and_customer_across_sheets(
    tmp_path: Path,
) -> None:
    path = _save(_clean_workbook(), tmp_path, name="sheet.xlsx", a1_dimension=False)
    _assert_fixture_shape(path, a1_dimension=False)

    parsed = read_workbook(path)

    assert {v.variety for v in parsed.ventas} == {VARIETY}
    assert parsed.entradas[0].variety == VARIETY
    assert {v.raw_name for v in parsed.varieties} & {VARIETY}
    assert {v.customer for v in parsed.ventas} == {CUSTOMER}


def test_read_workbook_over_the_a1_variant_reads_the_same(tmp_path: Path) -> None:
    path = _save(_clean_workbook(), tmp_path, name="sheet.xlsx", a1_dimension=True)
    _assert_fixture_shape(path, a1_dimension=True)

    parsed = read_workbook(path)

    _assert_clean_parse(parsed)


# -- The CLI dry run and real run over the clean fixture ------------------------------


def test_dry_run_over_the_clean_export_finds_everything_and_reports_mismatches(
    client: TestClient, engine: Engine, clock: MovableClock, tmp_path: Path
) -> None:
    business_id = _business_id(engine)
    xlsx = _save(_clean_workbook(), tmp_path, name="sheet.xlsx", a1_dimension=False)
    _assert_fixture_shape(xlsx, a1_dimension=False)
    mapping = _save_mapping(tmp_path)
    out, err = io.StringIO(), io.StringIO()

    code = run(
        [str(business_id), str(xlsx), "--mapping", str(mapping), "--dry-run"],
        engine=engine,
        clock=clock,
        out=out,
        err=err,
    )

    report = out.getvalue()
    assert code == 0, report
    assert err.getvalue() == ""
    assert "column_missing" not in report
    assert "Sales: 2" in report
    assert "Writable: yes" in report
    assert report.count("Total mismatches") == 1
    assert report.count("\n  Ventas row") == 2  # both non-whole Totals are mismatches
    assert client.get("/varieties").json() == []


def test_real_run_over_the_clean_export_writes_the_numeric_names(
    client: TestClient, engine: Engine, clock: MovableClock, tmp_path: Path
) -> None:
    business_id = _business_id(engine)
    xlsx = _save(_clean_workbook(), tmp_path, name="sheet.xlsx", a1_dimension=False)
    _assert_fixture_shape(xlsx, a1_dimension=False)
    mapping = _save_mapping(tmp_path)
    out, err = io.StringIO(), io.StringIO()

    code = run(
        [str(business_id), str(xlsx), "--mapping", str(mapping)],
        engine=engine,
        clock=clock,
        out=out,
        err=err,
    )

    assert code == 0, err.getvalue()
    assert err.getvalue() == ""

    varieties = client.get("/varieties").json()
    assert {v["name"] for v in varieties} == {VARIETY, "Blue Dream"}

    customers = client.get("/customers").json()
    assert [c["name"] for c in customers] == [CUSTOMER]


# -- The CLI dry run and real run over the problem variant ----------------------------


def test_dry_run_over_the_problem_export_lists_every_problem_in_one_pass(
    client: TestClient, engine: Engine, clock: MovableClock, tmp_path: Path
) -> None:
    business_id = _business_id(engine)
    xlsx = _save(_problem_workbook(), tmp_path, name="sheet.xlsx", a1_dimension=False)
    _assert_fixture_shape(xlsx, a1_dimension=False)
    mapping = _save_mapping(tmp_path)
    out, err = io.StringIO(), io.StringIO()

    code = run(
        [str(business_id), str(xlsx), "--mapping", str(mapping), "--dry-run"],
        engine=engine,
        clock=clock,
        out=out,
        err=err,
    )

    report = out.getvalue()
    assert code == 0, report
    assert "catalog.indica_share_invalid" in report
    assert "import.origin_unknown" in report
    assert report.count("text_invalid") == 2  # the boolean Nota and the 0.0 Revisar
    assert report.count("text_is_date") == 4  # the time and the three dates
    assert "re-type the cell as text" in report
    assert "apostrophe" in report
    # The hint never shows the cell's own value.
    assert "2026-10-12" not in report
    assert "2026-10-13" not in report
    assert "2026-10-14" not in report
    assert "12:30" not in report
    assert "Left out as unfinished" not in report
    assert "Writable: no" in report


def test_real_run_over_the_problem_export_writes_nothing(
    client: TestClient, engine: Engine, clock: MovableClock, tmp_path: Path
) -> None:
    business_id = _business_id(engine)
    xlsx = _save(_problem_workbook(), tmp_path, name="sheet.xlsx", a1_dimension=False)
    _assert_fixture_shape(xlsx, a1_dimension=False)
    mapping = _save_mapping(tmp_path)
    out, err = io.StringIO(), io.StringIO()

    code = run(
        [str(business_id), str(xlsx), "--mapping", str(mapping)],
        engine=engine,
        clock=clock,
        out=out,
        err=err,
    )

    assert code == 1
    assert err.getvalue() == ""
    assert client.get("/varieties").json() == []
    assert client.get("/customers").json() == []
    assert client.get("/stock").json()["varieties"] == []
