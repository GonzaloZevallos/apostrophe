"""The workbook the sample app hands to the real `scripts/import-spreadsheet.sh` (GH-274).

    build(path, today)

writes one small .xlsx for an empty business: one variety priced from `today`, one
finished "Stock inicial" entry dated `today`, empty Ventas and Gastos sheets with their
titles, and the categories and Catálogo cells the import needs. It imports no `app` module
and reads no clock: `today` is the one date the sample app was given.

The layout is hard-coded here on purpose, so that a change in the import is noticed by the
sample's dry run and fixed in this file, never in `app/`. Its sources are
`docs/product/import.md` and `app/spreadsheet_import/adapters/layout.py`:

    Ajustes   A7:A16 varieties (F7:F16 thresholds), A21:A28 categories,
              price history from row 37: variety, "Rige desde", list price
    Catálogo  rows 6-15: variety, shown, emoji, % índica, aroma; B18 header, B19 footer
    Entradas  row 1 titles, data from row 2, "Revisar" blank on a finished row
    Ventas    row 1 titles, with "Total"; Gastos: row 1 titles, with "Revisar"
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook

VARIETY = "Variedad importada"
CATEGORY = "Packaging"

VENTAS_TITLES = (
    "Fecha",
    "Variedad",
    "Gramos",
    "Pago",
    "Cliente",
    "Tipo",
    "Envío cobrado",
    "Precio especial (total)",
    "Nota",
    "Total",
)
ENTRADAS_TITLES = ("Fecha", "Variedad", "Gramos", "Origen", "Costo total", "Nota", "Revisar")
GASTOS_TITLES = ("Fecha", "Categoría", "Monto", "Detalle", "Revisar")


def build(path: Path, today: date) -> None:
    workbook = Workbook()
    ajustes = workbook.active
    assert ajustes is not None
    ajustes.title = "Ajustes"
    ajustes["A7"] = VARIETY
    ajustes["A21"] = CATEGORY
    ajustes["A37"] = VARIETY
    ajustes["B37"] = today
    ajustes["C37"] = 10_000

    catalogo = workbook.create_sheet("Catálogo")
    for column, value in zip("ABCDE", (VARIETY, "Sí", "🌿", 50, "Herbal"), strict=True):
        catalogo[f"{column}6"] = value
    catalogo["B18"] = "Catálogo"
    catalogo["B19"] = "Pedidos por mensaje"

    entradas = workbook.create_sheet("Entradas")
    entradas.append(ENTRADAS_TITLES)
    entradas.append((today, VARIETY, 20, "Stock inicial", None, None, None))

    workbook.create_sheet("Ventas").append(VENTAS_TITLES)
    workbook.create_sheet("Gastos").append(GASTOS_TITLES)
    workbook.save(path)
